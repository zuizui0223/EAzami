#!/usr/bin/env python3
"""Build an anonymous, line-numbered JEB V9.5 DOCX and a separate title page."""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "docs" / "chapter2" / "MANUSCRIPT_JEB_V9_5_ORGANIZATION_SPINE.md"


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--fig-dir", type=Path, required=True)
    p.add_argument("--output-dir", type=Path, required=True)
    return p.parse_args()


def clean_inline(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = text.replace("`", "")
    text = text.replace("\\(", "").replace("\\)", "")
    return text.strip()


def add_markdown_runs(paragraph, text: str) -> None:
    """Render the small inline-markdown subset used by the manuscript."""
    text = text.replace("`", "").replace("\\(", "").replace("\\)", "")
    # **bold**, *italic*; no nested emphasis is used in the manuscript.
    pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2]); r.bold = True
        else:
            r = paragraph.add_run(token[1:-1]); r.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def set_doc_defaults(doc: Document, anonymous: bool = True) -> None:
    sec = doc.sections[0]
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 2; normal.paragraph_format.space_after = Pt(0)

    for name, size in [("Title", 16), ("Heading 1", 14), ("Heading 2", 12)]:
        st = styles[name]
        st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    core = doc.core_properties
    core.title = "Repeated reassembly of a complex reproductive phenotype across unequal evolutionary depths in a young thistle radiation"
    core.subject = "Journal of Evolutionary Biology research article"
    core.author = "Anonymous" if anonymous else ""
    core.last_modified_by = ""; core.comments = ""


def add_line_numbers(section) -> None:
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:lnNumType")):
        sectPr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1"); ln.set(qn("w:start"), "1"); ln.set(qn("w:restart"), "continuous")
    sectPr.append(ln)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def set_image_alt(inline_shape, alt: str) -> None:
    docPr = inline_shape._inline.docPr
    docPr.set("descr", alt); docPr.set("title", alt[:120])


def add_figure(doc: Document, fig_n: int, fig_dir: Path, legend: str, alt: str) -> None:
    stems = {1:"figure1_repeated_components.png",2:"figure2_depth_ordering.png",3:"figure3_shared_localization.png",4:"figure4_representation_dependent_ecology.png"}
    path = fig_dir / stems[fig_n]
    if not path.exists(): raise FileNotFoundError(path)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(6.35)); set_image_alt(shape, alt)
    cap = doc.add_paragraph(); cap.paragraph_format.line_spacing = 1.15
    cap.paragraph_format.keep_together = True
    cap.add_run(f"Figure {fig_n}. ").bold = True
    add_markdown_runs(cap, legend)


def extract_figure_legends(text: str):
    if "# Figure legends" not in text or "# References" not in text: raise ValueError("expected Figure legends and References sections")
    main, rest = text.split("# Figure legends", 1); figs, tail = rest.split("# References", 1)
    legends, alts = {}, {}
    pattern = re.compile(r"\*\*Figure (\d+)\. (.*?)\*\*\s*(.*?)(?=\n\n\*\*Figure \d+\.|\Z)", re.S)
    for m in pattern.finditer(figs.strip()):
        n = int(m.group(1)); title = m.group(2).strip(); body = m.group(3).strip()
        parts = re.split(r"\*\*Alt text:\*\*", body, maxsplit=1)
        desc = parts[0].strip(); alt = clean_inline(parts[1]) if len(parts) == 2 else clean_inline(title)
        legends[n] = f"{title} {desc}".strip(); alts[n] = alt
    return main.strip(), "# References\n" + tail.strip(), legends, alts


def tokenize_blocks(text: str):
    lines=text.splitlines(); blocks=[]; buf=[]; in_math=False; mathbuf=[]
    for line in lines:
        s=line.rstrip()
        if s.strip()=="\\[":
            if buf: blocks.append(("p"," ".join(x.strip() for x in buf).strip())); buf=[]
            in_math=True; mathbuf=[]; continue
        if s.strip()=="\\]": blocks.append(("math"," ".join(mathbuf).strip())); in_math=False; mathbuf=[]; continue
        if in_math: mathbuf.append(s.strip()); continue
        if not s.strip():
            if buf: blocks.append(("p"," ".join(x.strip() for x in buf).strip())); buf=[]
            continue
        if s.startswith("# ") or s.startswith("## "):
            if buf: blocks.append(("p"," ".join(x.strip() for x in buf).strip())); buf=[]
            level=1 if s.startswith("# ") else 2; blocks.append((f"h{level}",s.lstrip("# ").strip()))
        elif s.startswith("- "):
            if buf: blocks.append(("p"," ".join(x.strip() for x in buf).strip())); buf=[]
            blocks.append(("bullet",s[2:].strip()))
        else: buf.append(s)
    if buf: blocks.append(("p"," ".join(x.strip() for x in buf).strip()))
    return blocks


def build_main(fig_dir: Path, out: Path) -> None:
    raw=MANUSCRIPT.read_text(encoding="utf-8")
    main,refs_tail,legends,alts=extract_figure_legends(raw)
    main=re.sub(r"\n\*\*Target journal:.*?\n\*\*Running title:.*?\n","\n",main,flags=re.S)
    refs_tail=refs_tail.split("## Submission-preparation notes",1)[0].strip()
    text=main+"\n\n"+refs_tail

    doc=Document(); set_doc_defaults(doc,anonymous=True); add_line_numbers(doc.sections[0]); add_page_number(doc.sections[0].footer.paragraphs[0])
    inserted=set()
    for kind,content in tokenize_blocks(text):
        if not content: continue
        if kind=="h1":
            if content==clean_inline(raw.splitlines()[0].lstrip("# ")):
                p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_markdown_runs(p,content)
            else:
                p=doc.add_paragraph(style="Heading 1"); add_markdown_runs(p,content)
        elif kind=="h2":
            p=doc.add_paragraph(style="Heading 2"); add_markdown_runs(p,content)
        elif kind=="bullet":
            p=doc.add_paragraph(style="List Bullet"); add_markdown_runs(p,content)
        elif kind=="math":
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(clean_inline(content)); r.italic=True
        else:
            p=doc.add_paragraph(); add_markdown_runs(p,content)
            for n in range(1,5):
                if n not in inserted and re.search(rf"\(Figure {n}\)",content): add_figure(doc,n,fig_dir,legends[n],alts[n]); inserted.add(n)
    if inserted!={1,2,3,4}: raise AssertionError(f"not all figures embedded: {sorted(inserted)}")
    out.parent.mkdir(parents=True,exist_ok=True); doc.save(out)


def build_title_page(out: Path) -> None:
    doc=Document(); set_doc_defaults(doc,anonymous=False)
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Repeated reassembly of a complex reproductive phenotype across unequal evolutionary depths in a young thistle radiation").bold=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Research Article - Journal of Evolutionary Biology")
    doc.add_paragraph(); p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Author names, affiliations, corresponding-author details and acknowledgements to be added at submission freeze.").italic=True
    doc.add_paragraph(); p=doc.add_paragraph(); p.add_run("Running title: ").bold=True; p.add_run("Unequal-depth capitulum reassembly")
    p=doc.add_paragraph(); p.add_run("Data/code archive: ").bold=True; p.add_run("Immutable DOI to be inserted after manuscript-evidence synchronization.")
    out.parent.mkdir(parents=True,exist_ok=True); doc.save(out)


def main() -> None:
    a=parse_args(); a.output_dir.mkdir(parents=True,exist_ok=True)
    build_main(a.fig_dir,a.output_dir/"EAzami_JEB_V9_5_anonymous_main.docx")
    build_title_page(a.output_dir/"EAzami_JEB_V9_5_title_page.docx")
    print(a.output_dir/"EAzami_JEB_V9_5_anonymous_main.docx"); print(a.output_dir/"EAzami_JEB_V9_5_title_page.docx")


if __name__=="__main__": main()
