#!/usr/bin/env python3
"""Build the active JEB V10 double-anonymous submission package."""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document

import build_chapter2_jeb_docx_v1 as legacy
import build_chapter2_jeb_docx_v9_6 as helpers

ROOT = Path(__file__).resolve().parents[1]
CH = ROOT / "docs" / "chapter2"
OUT = CH / "submission_package_v10"

MAN = CH / "MANUSCRIPT_JEB_V10_FUNCTIONAL_ASSEMBLY_ORIGIN_DECOUPLING.md"
TITLE = CH / "JEB_TITLE_PAGE_V10.md"
SI = CH / "JEB_SUPPORTING_INFORMATION_V5.md"
COVER = CH / "JEB_COVER_LETTER_V10.md"

FIGURE_SPECS = [
    {
        "number": 1,
        "section": "Three capitulum components repeatedly change within one young radiation",
        "filename": "figure1_v7_diversity_context.png",
        "caption": "Figure 1. Three capitulum components repeatedly change within one young radiation. Thirty-six of 38 sampled Japanese taxon concepts occur within the dominant radiation. Orientation requires four to six minimum state changes across the topology ensemble, phyllary posture exactly three and involucre stickiness exactly five.",
        "alt": "Alternative text: The young Japanese radiation contains repeated state changes in all three capitulum components; counts are lower bounds and are not labelled as adaptive origins.",
    },
    {
        "number": 2,
        "section": "The capitulum is assembled mosaically through evolutionary time",
        "filename": "figure2_v7_mosaic_depth.png",
        "caption": "Figure 2. Repeated capitulum assembly is historically mosaic. Paired same-topology relative-depth ordering shows phyllary posture deeper-permissive than stickiness in 1000/1000 topologies, phyllary deeper than orientation in 993/1000 and orientation deeper than stickiness in 905/1000. Zero of three trait pairs passes the robust shared-transition-localization rule.",
        "alt": "Alternative text: Depth ordering and branch-localization diagnostics show that the three components repeatedly change at different historical layers and do not share one synchronized branch history.",
    },
    {
        "number": 3,
        "section": "Repeated components map to plausible functional interfaces",
        "filename": "figure3_v10_functional_interfaces.png",
        "caption": "Figure 3. Repeated components map to distinct candidate functional interfaces. Orientation is linked to abiotic reproductive exposure/timing, phyllary architecture to mechanical antagonist access and stickiness to context-dependent arthropod filtering. The Cirsium reproductive-herbivory meta-analysis gives RR = 2.674 for seed output under reduced versus ambient insect herbivory, providing fitness-pressure context rather than a trait-specific adaptive effect.",
        "alt": "Alternative text: Three functional-interface columns summarize direct manipulation evidence and claim boundaries, with a separate meta-analytic fitness-pressure panel for reproductive herbivory.",
    },
    {
        "number": 4,
        "section": "Orientation transitions track a present ecological regime",
        "filename": "figure4_v10_orientation_transition_regime.png",
        "caption": "Figure 4. Reconstructed orientation transitions track a present two-axis ecological regime. The fixed U-to-D BIO15-up/BIO1-down composite ranks 16/792, 19/1716 and 4/126 across the three occurrence thresholds; the strict bidirectional-floor rank is 3/126. Direction survives all single-taxon deletions and declared geography/internal-edge stresses.",
        "alt": "Alternative text: Exact finite-map ranks and robustness summaries show structured transition-level orientation correspondence with higher precipitation seasonality and lower temperature without interpreting the result as selection.",
    },
    {
        "number": 5,
        "section": "The present orientation regime fails as a model of historical origin",
        "filename": "figure5_v10_origin_trigger_falsification.png",
        "caption": "Figure 5. The present orientation regime does not identify the historical trigger. The frozen historical-persistence criterion required at least 75% matching chronologies in each of all four palaeolocation regions; no region approaches that threshold. The BIO15-up/BIO1-down regime matches 99/376 deterministic chronology-by-palaeolocation scenarios for the sole bounded U-to-D event. At the central 0.79-0.74 Ma chronology BIO15 decreases, rather than increases, in all four regions. Broader diagnostics yield 0/324 robust climate classes and 0/21 robust sea-level classes.",
        "alt": "Alternative text: Historical scenario match fractions and central-chronology sign changes show failure of the present ecological regime as a persistent origin model, followed by the broader climate and sea-level identifiability ceiling.",
    },
]


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--output-dir", type=Path, default=OUT)
    p.add_argument("--figure-dir", type=Path, required=True)
    return p.parse_args()


def words(text: str) -> int:
    return len(re.findall(r"\b[\w'’*-]+\b", text))


def manuscript_counts() -> tuple[int, int]:
    text = MAN.read_text(encoding="utf-8")
    abstract = text.split("## Abstract\n\n", 1)[1].split("\n\n**Keywords:**", 1)[0]
    body = text.split("# References", 1)[0]
    return words(body), words(abstract)


def prepare_main_markdown(output_dir: Path) -> Path:
    text = MAN.read_text(encoding="utf-8")
    if "# Figure legends" not in text or "# References" not in text:
        raise RuntimeError("Could not isolate Figure legends / References")
    before, rest = text.split("# Figure legends", 1)
    _, references = rest.split("# References", 1)
    rendered = before.rstrip() + "\n\n# References\n" + references.lstrip()
    rendered = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", rendered)
    rendered = re.sub(r"\\\((.*?)\\\)", r"\1", rendered)
    rendered = rendered.replace("\\[", "").replace("\\]", "")
    rendered = rendered.replace("D=(N-d)/(N-1).", "D = (N - d)/(N - 1).")
    rendered = rendered.replace("D=1", "D = 1")
    path = output_dir / "_render_main_v10.md"
    path.write_text(rendered, encoding="utf-8")
    return path


def prepare_title_markdown(output_dir: Path) -> Path:
    main_count, abstract_count = manuscript_counts()
    text = TITLE.read_text(encoding="utf-8")
    text = text.split("## Submission check", 1)[0].rstrip() + "\n"
    text = text.replace("{{MAIN_WORD_COUNT}}", str(main_count))
    text = text.replace("{{ABSTRACT_WORD_COUNT}}", str(abstract_count))
    path = output_dir / "_render_title_v10.md"
    path.write_text(text, encoding="utf-8")
    return path


def build_main(output_dir: Path, figure_dir: Path) -> Path:
    source = prepare_main_markdown(output_dir)
    doc = Document()
    legacy.configure_document(doc, running_header="Rebuilding the Cirsium capitulum", line_numbers=True)
    helpers.blank_identifying_metadata(doc)
    legacy.render_markdown(
        doc,
        source,
        skip_metadata_prefixes=("**Target journal:**", "**Status:**", "**Running title:**"),
    )
    helpers.assert_anonymous_visible_text(doc)
    for spec in FIGURE_SPECS:
        helpers.add_figure_after(
            helpers.first_body_paragraph_after_heading(doc, spec["section"]),
            figure_dir / spec["filename"],
            spec["caption"],
            spec["alt"],
        )
    path = output_dir / "Chapter2_JEB_Anonymous_Manuscript_V10.docx"
    legacy.save_document(doc, path)
    helpers.scrub_package(path)
    source.unlink(missing_ok=True)
    return path


def build_title_page(output_dir: Path) -> Path:
    source = prepare_title_markdown(output_dir)
    doc = Document()
    legacy.configure_document(doc, running_header="", line_numbers=False)
    helpers.blank_identifying_metadata(doc)
    legacy.render_markdown(doc, source)
    path = output_dir / "Chapter2_JEB_Title_Page_V10.docx"
    legacy.save_document(doc, path)
    helpers.scrub_package(path)
    source.unlink(missing_ok=True)
    return path


def build_supporting(output_dir: Path) -> Path:
    doc = Document()
    legacy.configure_document(doc, running_header="Supporting Information", line_numbers=True)
    helpers.blank_identifying_metadata(doc)
    legacy.render_markdown(doc, SI, stop_heading="Supporting-information audit checklist")
    path = output_dir / "Chapter2_JEB_Supporting_Information_V10.docx"
    legacy.save_document(doc, path)
    helpers.scrub_package(path)
    return path


def build_cover_letter(output_dir: Path) -> Path:
    doc = Document()
    legacy.configure_document(doc, running_header="", line_numbers=False)
    helpers.blank_identifying_metadata(doc)
    legacy.render_markdown(doc, COVER)
    path = output_dir / "Chapter2_JEB_Cover_Letter_V10.docx"
    legacy.save_document(doc, path)
    helpers.scrub_package(path)
    return path


def main():
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    outputs = [
        build_main(args.output_dir, args.figure_dir),
        build_title_page(args.output_dir),
        build_supporting(args.output_dir),
        build_cover_letter(args.output_dir),
    ]
    for path in outputs:
        if not path.exists() or path.stat().st_size < 1000:
            raise RuntimeError(f"Document build failed or unexpectedly small: {path}")
        print(path, path.stat().st_size)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
