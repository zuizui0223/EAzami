#!/usr/bin/env python3
"""Build the active JEB V9.6 double-anonymous submission package."""
from __future__ import annotations

import argparse
import re
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from lxml import etree

import build_chapter2_jeb_docx_v1 as legacy

ROOT = Path(__file__).resolve().parents[1]
CH = ROOT / "docs" / "chapter2"
OUT = CH / "submission_package_v9_6"

MAN = CH / "MANUSCRIPT_JEB_V9_6_SCALE_CONDITIONED_ECOLOGY.md"
TITLE = CH / "JEB_TITLE_PAGE_V9_6.md"
SI = CH / "JEB_SUPPORTING_INFORMATION_V4.md"
COVER = CH / "JEB_COVER_LETTER_V9_6.md"

FIGURE_SPECS = [
    {
        "number": 1,
        "section": "Repeated change occurs within one young radiation",
        "filename": "figure1_v9_6_repeated_components.png",
        "caption": "Figure 1. Repeated component differentiation occurs within one young radiation. Thirty-six of 38 sampled Japanese taxon concepts occur in the dominant radiation. Orientation, phyllary posture and involucre stickiness each require repeated minimum state changes under the admitted topology ensemble. Missing and source-conflicting trait states remain explicit.",
        "alt": "Alternative text: Three-panel figure showing 36 of 38 sampled Japanese taxon concepts in one dominant radiation, minimum-change ranges for orientation, phyllary posture and stickiness, and an authority-backed trait-state matrix with unknown and conflicting states retained.",
    },
    {
        "number": 2,
        "section": "Repeated histories are stratified across evolutionary depth",
        "filename": "figure2_v9_6_unequal_depth.png",
        "caption": "Figure 2. Repeated histories occupy unequal evolutionary depths. Relative-depth envelopes, paired same-topology ordering and coverage-matched sensitivity show a robust central ordering in which phyllary posture is deeper-permissive and stickiness shallower, while strict deepest tails overlap after coverage matching. Relative lineage depth is topology-only and is not calendar time or an evolutionary rate.",
        "alt": "Alternative text: Four-panel figure showing relative lineage-depth envelopes for three capitulum traits, paired topology ordering of 1000 of 1000, 993 of 1000 and 905 of 1000, coverage-matched central versus strict-tail results, and a summary that central depth ordering is robust while deepest tails overlap.",
    },
    {
        "number": 3,
        "section": "Component changes are not repeatedly synchronized on the same branches",
        "filename": "figure3_v9_6_shared_localization.png",
        "caption": "Figure 3. Component changes do not repeatedly synchronize on the same branches. Branch-length-aware transition-posterior overlap and equal-branch topology sensitivities disagree or include non-positive tails for all three trait pairs; zero of three pairs passes the robust shared-localization rule.",
        "alt": "Alternative text: Two-panel figure showing branch-length-aware point estimates and equal-branch q05-to-q95 transition-localization intervals for orientation-phyllary, orientation-stickiness and phyllary-stickiness, plus a summary box stating that zero of three trait pairs passes the robust shared-localization rule.",
    },
    {
        "number": 4,
        "section": "Orientation ecology is not scale invariant",
        "filename": "figure4_v9_6_scale_conditioned_ecology.png",
        "caption": "Figure 4. Orientation-environment correspondence changes with biological scale and phenotype representation. Static common9 state separation is weak or resolution-limited across the three traits; Azami among- and within-taxon image-angle associations differ from EAzami transition-level BIO1/BIO15 directions; the fixed transition-regime composite is exceptional under exact finite-map ranks; and the preregistered external confirmation stops before climate extraction because strict occurrence QC leaves only one upward taxon.",
        "alt": "Alternative text: Four-panel figure showing common nine-variable static ecology for three traits, a categorical cross-scale orientation matrix for BIO12 BIO1 and BIO15, exact transition-regime finite-map ranks across three occurrence thresholds, and a held-out confirmation flow that ends at zero downward and one upward eligible taxon with climate outcomes unopened.",
    },
]


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--output-dir", type=Path, default=OUT)
    p.add_argument("--figure-dir", type=Path, required=True)
    return p.parse_args()


def words(text: str) -> int:
    return len(re.findall(r"\b[\w'’*-]+\b", text))


def manuscript_counts() -> tuple[int, int]:
    text = MAN.read_text(encoding="utf-8")
    abstract = text.split("## Abstract\n\n", 1)[1].split("\n\n**Keywords:**", 1)[0]
    body = text.split("# References", 1)[0]
    return words(body), words(abstract)


def blank_identifying_metadata(doc: Document) -> None:
    props = doc.core_properties
    for attr in ("author", "last_modified_by", "comments", "category", "subject", "keywords"):
        setattr(props, attr, "")


def paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def add_figure_after(paragraph: Paragraph, image_path: Path, caption: str, alt: str) -> None:
    if not image_path.exists() or image_path.stat().st_size < 3000:
        raise RuntimeError(f"Missing or unexpectedly small figure: {image_path}")
    image_par = paragraph_after(paragraph)
    image_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = image_par.add_run().add_picture(str(image_path), width=Inches(6.20))
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", caption.split(".", 1)[0])
    caption_par = paragraph_after(image_par)
    try:
        caption_par.style = "Caption"
    except KeyError:
        pass
    caption_par.add_run(caption)


def first_body_paragraph_after_heading(doc: Document, heading_text: str) -> Paragraph:
    pars = list(doc.paragraphs)
    hits = [i for i, p in enumerate(pars) if p.text.strip() == heading_text]
    if len(hits) != 1:
        raise RuntimeError(f"Expected one heading {heading_text!r}, found {len(hits)}")
    for p in pars[hits[0] + 1 :]:
        if not p.text.strip():
            continue
        if p.style.name.startswith("Heading"):
            raise RuntimeError(f"No body paragraph below heading {heading_text!r}")
        return p
    raise RuntimeError(f"No body paragraph below heading {heading_text!r}")


def prepare_main_markdown(output_dir: Path) -> Path:
    text = MAN.read_text(encoding="utf-8")
    if "# Figure legends" not in text or "# References" not in text:
        raise RuntimeError("Could not isolate Figure legends / References")
    before, rest = text.split("# Figure legends", 1)
    _, references = rest.split("# References", 1)
    rendered = before.rstrip() + "\n\n# References\n" + references.lstrip()
    path = output_dir / "_render_main_v9_6.md"
    path.write_text(rendered, encoding="utf-8")
    return path


def prepare_title_markdown(output_dir: Path) -> Path:
    main_count, abstract_count = manuscript_counts()
    text = TITLE.read_text(encoding="utf-8")
    text = text.replace("{{MAIN_WORD_COUNT}}", str(main_count))
    text = text.replace("{{ABSTRACT_WORD_COUNT}}", str(abstract_count))
    path = output_dir / "_render_title_v9_6.md"
    path.write_text(text, encoding="utf-8")
    return path


def assert_anonymous_visible_text(doc: Document) -> None:
    visible = "\n".join(p.text for p in doc.paragraphs)
    low = visible.casefold()
    bad = []
    for token in (
        "github.com/zuizui0223",
        "zuizui0223",
        "[insert",
        "data availability statement",
        "acknowledgements",
        "conflict of interest",
    ):
        if token in low:
            bad.append(token)
    if bad:
        raise RuntimeError(f"Anonymous main retains identifying/title-page-only text: {bad}")


def scrub_package(path: Path) -> None:
    tmp = path.with_suffix(".scrubbed.docx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xml"):
                try:
                    root = etree.fromstring(data)
                    for elem in root.iter():
                        for key in list(elem.attrib):
                            if etree.QName(key).localname.startswith("rsid"):
                                del elem.attrib[key]
                    for elem in list(root.iter()):
                        if etree.QName(elem.tag).localname.startswith("rsid"):
                            parent = elem.getparent()
                            if parent is not None:
                                parent.remove(elem)
                    if item.filename == "docProps/core.xml":
                        ns = {
                            "dc": "http://purl.org/dc/elements/1.1/",
                            "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                        }
                        for xpath in ("//dc:creator", "//cp:lastModifiedBy"):
                            for node in root.xpath(xpath, namespaces=ns):
                                node.text = ""
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                except Exception:
                    pass
            zout.writestr(item, data)
    tmp.replace(path)


def build_main(output_dir: Path, figure_dir: Path) -> Path:
    source = prepare_main_markdown(output_dir)
    doc = Document()
    legacy.configure_document(doc, running_header="Unequal-depth capitulum reassembly", line_numbers=True)
    blank_identifying_metadata(doc)
    legacy.render_markdown(
        doc,
        source,
        skip_metadata_prefixes=("**Target journal:**", "**Status:**", "**Running title:**"),
    )
    assert_anonymous_visible_text(doc)
    for spec in FIGURE_SPECS:
        add_figure_after(
            first_body_paragraph_after_heading(doc, spec["section"]),
            figure_dir / spec["filename"],
            spec["caption"],
            spec["alt"],
        )
    path = output_dir / "Chapter2_JEB_Anonymous_Manuscript_V9_6.docx"
    legacy.save_document(doc, path)
    scrub_package(path)
    source.unlink(missing_ok=True)
    return path


def build_title_page(output_dir: Path) -> Path:
    source = prepare_title_markdown(output_dir)
    doc = Document()
    legacy.configure_document(doc, running_header="", line_numbers=False)
    blank_identifying_metadata(doc)
    legacy.render_markdown(doc, source)
    path = output_dir / "Chapter2_JEB_Title_Page_V9_6.docx"
    legacy.save_document(doc, path)
    scrub_package(path)
    source.unlink(missing_ok=True)
    return path


def build_supporting(output_dir: Path) -> Path:
    doc = Document()
    legacy.configure_document(doc, running_header="Supporting Information", line_numbers=True)
    blank_identifying_metadata(doc)
    legacy.render_markdown(doc, SI)
    path = output_dir / "Chapter2_JEB_Supporting_Information_V9_6.docx"
    legacy.save_document(doc, path)
    scrub_package(path)
    return path


def build_cover_letter(output_dir: Path) -> Path:
    doc = Document()
    legacy.configure_document(doc, running_header="", line_numbers=False)
    blank_identifying_metadata(doc)
    legacy.render_markdown(doc, COVER)
    path = output_dir / "Chapter2_JEB_Cover_Letter_V9_6.docx"
    legacy.save_document(doc, path)
    scrub_package(path)
    return path


def main():
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    outputs = [
        build_main(args.output_dir, args.figure_dir),
        build_title_page(args.output_dir),
        build_supporting(args.output_dir),
        build_cover_letter(args.output_dir),
    ]
    for path in outputs:
        if not path.exists() or path.stat().st_size < 1000:
            raise RuntimeError(f"Document build failed or unexpectedly small: {path}")
        print(path, path.stat().st_size)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
